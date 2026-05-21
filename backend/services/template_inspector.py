from pathlib import Path
import re
from typing import Dict, Set
from docx import Document
try:
    from docxtpl import DocxTemplate
except Exception:
    DocxTemplate = None

PLACEHOLDER_RE = re.compile(r"\{\{\s*([^\}]+)\s*\}\}")

def extract_placeholders_from_docx(path: Path) -> Set[str]:
    placeholders: Set[str] = set()
    try:
        # First try docxtpl extractor if available (detects jinja variables)
        if DocxTemplate is not None:
            try:
                tpl = DocxTemplate(str(path))
                vars = tpl.get_undeclared_template_variables()
                for v in vars:
                    placeholders.add(str(v))
                # if we found vars, return early
                if placeholders:
                    return placeholders
            except Exception:
                pass

        doc = Document(path)
        # paragraphs
        for p in doc.paragraphs:
            for m in PLACEHOLDER_RE.findall(p.text):
                placeholders.add(m.strip())
        # tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for m in PLACEHOLDER_RE.findall(cell.text):
                        placeholders.add(m.strip())
        # fallback: search for other patterns like $var or <<var>>
        EXTRA_RE = re.compile(r"\$\{?([A-Za-z0-9_\-]+)\}?|<<\s*([A-Za-z0-9_\-]+)\s*>>")
        for p in doc.paragraphs:
            for m in EXTRA_RE.findall(p.text):
                for g in m:
                    if g:
                        placeholders.add(g)
    except Exception:
        # If reading fails, return empty set
        pass
    return placeholders


def scan_templates(templates_dir: Path) -> Dict[str, Set[str]]:
    mapping: Dict[str, Set[str]] = {}
    for p in sorted(templates_dir.glob('*.docx')):
        mapping[p.name] = extract_placeholders_from_docx(p)
    return mapping
